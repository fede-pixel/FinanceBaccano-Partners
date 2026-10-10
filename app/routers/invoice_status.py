"""Situazione fatture: da pagare, da incassare, pagate (con il movimento che lo prova)."""
import io
from datetime import date

from fastapi import APIRouter, Depends, Form, HTTPException, Request
from fastapi.responses import Response
from sqlalchemy.orm import Session

from app.db import get_session
from app.models import PlannedItem, Transaction
from app.services import invoice_match, reconcile
from app.services.util import get_setting, parse_date, set_setting
from app.web import redirect, render

router = APIRouter()

VIEWS = {
    "pagare": ("passiva", ("aperta",), "Da pagare"),
    "incassare": ("attiva", ("aperta",), "Da incassare"),
    "presunte": (None, ("presunta",), "Pagate senza movimento"),
    "pagate": (None, ("pagata", "compensata"), "Pagate e incassate"),
    "contenzioso": (None, ("contenzioso", "perdita"), "Contenzioso e perdite"),
}


def _before(rows: list[invoice_match.InvoiceStatus], start: date | None) -> list[invoice_match.InvoiceStatus]:
    """Fatture aperte con data anteriore al periodo coperto dai movimenti: il pagamento può essere avvenuto prima."""
    return [s for s in rows if start and s.state == "aperta" and s.invoice.date < start]


def _rows(session: Session, vista: str, q: str = "", lotto: str = "", anteriori: bool = False) -> list[invoice_match.InvoiceStatus]:
    direction, states, _ = VIEWS[vista]
    rows = [s for s in invoice_match.invoice_statuses(session, direction) if s.state in states]
    if lotto:
        rows = [s for s in rows if any(it.match_batch == lotto for it in s.items)]
    if anteriori:
        rows = _before(rows, invoice_match.coverage(session).start)
    if q:
        ql = q.lower()
        rows = [s for s in rows if ql in s.invoice.counterparty.lower() or ql in s.invoice.number.lower()]
    if states == ("aperta",):
        rows.sort(key=lambda s: (s.due or date.max, s.invoice.counterparty))
    else:
        rows.sort(key=lambda s: (s.paid_date or s.invoice.date), reverse=True)
    return rows


@router.get("/fatture/situazione")
def situation(request: Request, vista: str = "pagare", q: str = "", lotto: str = "", anteriori: int = 0,
              session: Session = Depends(get_session)):
    if vista not in VIEWS:
        vista = "pagare"
    result = invoice_match.run(session)  # gli abbinamenti certi si fanno da soli a ogni apertura
    session.commit()
    all_rows = invoice_match.invoice_statuses(session)
    today = date.today()

    cov = invoice_match.coverage(session)

    def kpi(direction, state="aperta"):
        rows = [s for s in all_rows if s.state == state and (direction is None or s.invoice.direction == direction)]
        before = _before(rows, cov.start)
        return {
            "before_count": len(before),
            "before_total": sum(s.open_cents for s in before),
            "count": len(rows),
            "total": sum(s.open_cents if state == "aperta" else s.paid_cents for s in rows),
            "overdue": sum(s.open_cents for s in rows if s.due and s.due < today),
            "overdue_count": sum(1 for s in rows if s.due and s.due < today),
        }

    msg = request.query_params.get("msg", "")
    if result.matched or result.compensated:
        auto = f"Controllo incrociato: {result.matched} scadenze abbinate ai movimenti"
        if result.compensated:
            auto += f", {result.compensated} note di credito compensate"
        msg = f"{msg} · {auto}" if msg else auto
    last = get_setting(session, "last_match_batch").split("|")  # lotto|data e ora|descrizione dell'ultima azione in blocco
    last_batch = last[0] if invoice_match.batch_size(session, last[0]) else ""
    rows = _rows(session, vista, q, lotto, bool(anteriori))
    return render(
        request, "invoice_status.html", vista=vista, views=VIEWS, q=q, rows=rows,
        diag=invoice_match.diagnose(session, rows) if vista in ("pagare", "incassare") else {},
        proposals=result.proposals, msg=msg, lotto=lotto, anteriori=anteriori, cov=cov,
        last_batch=last_batch, last_batch_at=last[1] if len(last) > 1 else "",
        last_batch_label=last[2] if len(last) > 2 else "Conferma in blocco",
        last_batch_size=invoice_match.batch_size(session, last_batch),
        k_pay=kpi("passiva"), k_inc=kpi("attiva"), k_pres=kpi(None, "presunta"), k_paid=kpi(None, "pagata"),
        k_disp={"count": sum(1 for x in all_rows if x.state in ("contenzioso", "perdita")),
                "disputed": sum(x.disputed_cents for x in all_rows), "lost": sum(x.lost_cents for x in all_rows)},
    )


@router.post("/fatture/situazione/conferma-tutte")
def confirm_all(session: Session = Depends(get_session)):
    from datetime import datetime

    batch, n = invoice_match.confirm_all(session)
    if n:
        set_setting(session, "last_match_batch", f"{batch}|{datetime.now():%d/%m/%Y %H:%M}|Conferma in blocco")
    session.commit()
    if not n:
        return redirect("/fatture/situazione", msg="Nessuna proposta da confermare")
    return redirect(f"/fatture/situazione?vista=pagate&lotto={batch}",
                    msg=f"Confermati in blocco {n} abbinamenti: qui sotto solo quelli, controllali. Puoi annullarli tutti con un clic.")


@router.post("/fatture/situazione/segna-anteriori")
def mark_before(vista: str = Form("incassare"), session: Session = Depends(get_session)):
    """Segna pagate/incassate (senza movimento) le fatture aperte anteriori ai movimenti importati."""
    from datetime import datetime

    if vista not in ("pagare", "incassare"):
        vista = "incassare"
    rows = _rows(session, vista, anteriori=True)
    batch, n = invoice_match.mark_presumed(session, rows)
    if n:
        label = "Fatture anteriori ai movimenti segnate " + ("incassate" if vista == "incassare" else "pagate")
        set_setting(session, "last_match_batch", f"{batch}|{datetime.now():%d/%m/%Y %H:%M}|{label}")
    session.commit()
    return redirect(f"/fatture/situazione?vista=presunte&lotto={batch}",
                    msg=f"{len(rows)} fatture anteriori ai movimenti segnate come saldate (senza movimento). Puoi annullare con un clic.")


@router.post("/fatture/situazione/annulla-lotto")
def undo_batch(lotto: str = Form(...), session: Session = Depends(get_session)):
    n = invoice_match.undo_batch(session, lotto)
    if get_setting(session, "last_match_batch").startswith(lotto):
        set_setting(session, "last_match_batch", "")
    session.commit()
    return redirect("/fatture/situazione", msg=f"Annullato: {n} scadenze tornate com'erano prima")


def _ids(raw: str) -> list[int]:
    return [int(i) for i in raw.split(",") if i.strip().isdigit()]


@router.post("/fatture/situazione/abbina")
def confirm(tx_ids: str = Form(...), item_ids: str = Form(...), session: Session = Depends(get_session)):
    txs = [session.get(Transaction, i) for i in _ids(tx_ids)]
    items = [session.get(PlannedItem, i) for i in _ids(item_ids)]
    if not txs or not items or any(x is None for x in txs + items):
        raise HTTPException(404, "Non trovato")
    for tx in txs:
        tx.no_auto_match = False  # abbinamento scelto da te
        if tx.planned_item_id:
            reconcile.unmatch(session, tx)
    for it in items:
        it.no_auto_match = False
    if len(txs) > 1:
        invoice_match.settle_installments(session, txs, items[0])
        msg = f"Abbinato: fattura saldata da {len(txs)} acconti"
    else:
        invoice_match.settle(session, txs[0], items)
        msg = f"Abbinato: {len(items)} fatture saldate dal movimento del {txs[0].date:%d/%m/%Y}"
    session.commit()
    return redirect("/fatture/situazione?vista=pagate", msg=msg)


@router.post("/fatture/situazione/scollega")
def unlink(tx_id: int = Form(...), back: str = Form("/fatture/situazione?vista=pagate"), session: Session = Depends(get_session)):
    tx = session.get(Transaction, tx_id)
    if tx is None:
        raise HTTPException(404, "Non trovato")
    reconcile.unmatch(session, tx)
    tx.no_auto_match = True  # non riabbinarlo da solo
    session.commit()
    return redirect(back, msg="Abbinamento annullato: le fatture tornano da saldare (l'app non lo rifarà da sola)")


@router.post("/fatture/situazione/segna")
def mark(item_ids: str = Form(...), action: str = Form("pagata"), paid_date: str = Form(""),
         back: str = Form("/fatture/situazione"), session: Session = Depends(get_session)):
    """Segna a mano come pagate (es. pagamento in contanti o da un conto non caricato) o riapre."""
    n = 0
    for i in item_ids.split(","):
        it = session.get(PlannedItem, int(i)) if i.strip().isdigit() else None
        if it is None:
            continue
        if action == "riapri":
            it.status, it.paid_date, it.settlement = "previsto", None, ""
            it.no_auto_match = True  # l'hai detto tu: non va considerata pagata in automatico
        else:
            it.status, it.settlement = "pagato", "manuale"
            it.paid_date = parse_date(paid_date) or date.today()
        n += 1
    session.commit()
    return redirect(back, msg="Fattura riaperta" if action == "riapri" else "Fattura segnata come pagata")


@router.post("/fatture/situazione/segna-selezionate")
async def mark_selected(request: Request, session: Session = Depends(get_session)):
    """Segna pagate/incassate le fatture spuntate (es. "tutte tranne una"), con lotto annullabile."""
    from datetime import datetime

    from app.models import Invoice

    form = await request.form()
    vista = form.get("vista", "pagare")
    ids = [int(v) for v in form.getlist("invoice_ids") if str(v).isdigit()]
    items = [it for i in ids if (inv := session.get(Invoice, i)) for it in inv.planned_items]
    batch, n = invoice_match.mark_items_paid(session, items, parse_date(form.get("paid_date", "")))
    if n:
        label = f"{len(ids)} fatture segnate " + ("incassate" if vista == "incassare" else "pagate") + " a mano"
        set_setting(session, "last_match_batch", f"{batch}|{datetime.now():%d/%m/%Y %H:%M}|{label}")
    session.commit()
    back = f"/fatture/situazione?vista={vista}"
    if not n:
        return redirect(back, err="Nessuna fattura selezionata")
    return redirect(back, msg=f"{len(ids)} fatture segnate come saldate. Puoi annullare con un clic dal riquadro in alto.")


@router.post("/fatture/situazione/parziale")
def partial(item_id: int = Form(...), amount: str = Form(...), paid_date: str = Form(""), vista: str = Form("pagare"),
            session: Session = Depends(get_session)):
    from app.services.util import fmt_eur, parse_amount

    it = session.get(PlannedItem, item_id)
    cents = parse_amount(amount)
    back = f"/fatture/situazione?vista={vista}"
    if it is None or cents is None:
        return redirect(back, err="Importo non valido")
    try:
        invoice_match.partial_payment(session, it, cents, parse_date(paid_date) or date.today())
    except ValueError as exc:
        return redirect(back, err=str(exc))
    session.commit()
    return redirect(back, msg=f"Pagamento parziale registrato: restano {fmt_eur(abs(it.amount_cents))} da saldare")


@router.post("/fatture/situazione/contenzioso")
def dispute(item_ids: str = Form(...), action: str = Form("contenzioso"), note: str = Form(""),
            back: str = Form("/fatture/situazione?vista=contenzioso"), session: Session = Depends(get_session)):
    """Contenzioso (legali), stralcio (credito perso) o riattivazione."""
    items = [it for i in _ids(item_ids) if (it := session.get(PlannedItem, i))]
    if action == "stralcia":
        n, msg = invoice_match.write_off(items), "Credito stralciato: registrato come perdita"
    elif action == "riattiva":
        n, msg = invoice_match.reactivate(items), "Credito riattivato: torna tra quelli da incassare"
    else:
        n, msg = invoice_match.set_dispute(items, note.strip()), ("Fattura in contenzioso: esce da incassi previsti, "
                                                                  "forecast e capacità di investimento")
    session.commit()
    return redirect(back if back.startswith("/") else "/fatture/situazione", msg=msg if n else "Nessuna modifica")


@router.get("/fatture/situazione.xlsx")
def export(vista: str = "pagare", session: Session = Depends(get_session)):
    import openpyxl
    from openpyxl.styles import Font

    if vista not in VIEWS:
        vista = "pagare"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = VIEWS[vista][2][:30]
    ws.append(["Tipo", "Data fattura", "Numero", "Documento", "Controparte", "P.IVA", "Scadenza", "Giorni di ritardo",
               "Totale fattura", "Da saldare", "Pagata il", "Movimento"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for s in _rows(session, vista):
        inv = s.invoice
        ws.append([
            "Ricevuta" if inv.direction == "passiva" else "Emessa", inv.date, inv.number, inv.doc_type, inv.counterparty,
            inv.counterparty_vat, s.due, s.overdue_days or None, inv.total_cents / 100, s.open_cents / 100 if s.open_cents else None,
            s.paid_date, "; ".join(f"{t.date:%d/%m/%Y} {t.account.name} {t.description[:60]}" for t in s.transactions),
        ])
    for col, w in zip("ABCDEFGHIJKL", (9, 12, 16, 8, 38, 14, 12, 10, 13, 13, 12, 60)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in (row[1], row[6], row[10]):
            c.number_format = "DD/MM/YYYY"
        for c in (row[8], row[9]):
            c.number_format = "#,##0.00"
    buf = io.BytesIO()
    wb.save(buf)
    name = f"fatture_{vista}_{date.today():%Y%m%d}.xlsx"
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.get("/controparte")
def counterparty_page(request: Request, q: str = "", session: Session = Depends(get_session)):
    from app.services import counterparty

    led = counterparty.build(session, q) if q.strip() else None
    return render(request, "counterparty.html", q=q, led=led)
