"""Import dei movimenti da file esportati dall'home banking (CSV, XLSX, XLS).

Ogni banca esporta un formato diverso: la mappatura delle colonne viene indovinata e
poi salvata sul conto, così dal secondo import in poi basta trascinare il file."""
import csv
import hashlib
import io
import json
import re
from pathlib import Path
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Account, BalanceSnapshot, Transaction
from app.services.util import parse_amount

DATE_FORMATS = [
    "%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%Y/%m/%d",
    "%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M",
    "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M:%S.%f", "%Y-%m-%dT%H:%M:%S.%fZ", "%Y-%m-%dT%H:%M:%SZ",
    "%d.%m.%y", "%d-%m-%y", "%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S",
]

# Mesi scritti a lettere (italiano, inglese, tedesco: Vivid è una banca tedesca)
_MONTH_NAMES = {
    "gen": 1, "gennaio": 1, "jan": 1, "january": 1, "januar": 1, "jän": 1,
    "feb": 2, "febbraio": 2, "february": 2, "februar": 2,
    "mar": 3, "marzo": 3, "march": 3, "mär": 3, "mrz": 3, "märz": 3,
    "apr": 4, "aprile": 4, "april": 4,
    "mag": 5, "maggio": 5, "may": 5, "mai": 5,
    "giu": 6, "giugno": 6, "jun": 6, "june": 6, "juni": 6,
    "lug": 7, "luglio": 7, "jul": 7, "july": 7, "juli": 7,
    "ago": 8, "agosto": 8, "aug": 8, "august": 8,
    "set": 9, "settembre": 9, "sep": 9, "sept": 9, "september": 9,
    "ott": 10, "ottobre": 10, "oct": 10, "october": 10, "okt": 10, "oktober": 10,
    "nov": 11, "novembre": 11, "november": 11,
    "dic": 12, "dicembre": 12, "dec": 12, "december": 12, "dez": 12, "dezember": 12,
}
_TEXT_DATE = re.compile(r"^(\d{1,2})[\s.\-/]+([a-zà-ü]{3,10})\.?[\s.\-/,]+(\d{2,4})\b", re.I)
_TEXT_DATE_US = re.compile(r"^([a-zà-ü]{3,10})\.?\s+(\d{1,2}),?\s+(\d{4})\b", re.I)
_LEADING_DATE = re.compile(r"^(\d{1,2}[/.\-]\d{1,2}[/.\-]\d{2,4}|\d{4}-\d{2}-\d{2})")


def _text_date(value: str) -> date | None:
    """'1 ott 2026', '01 Oct 2026', '1. Okt. 2026', 'Oct 1, 2026' -> date."""
    m = _TEXT_DATE.match(value)
    if m:
        day, mon, year = m.group(1), m.group(2), m.group(3)
    else:
        m = _TEXT_DATE_US.match(value)
        if not m:
            return None
        mon, day, year = m.group(1), m.group(2), m.group(3)
    month = _MONTH_NAMES.get(mon.lower())
    if not month:
        return None
    y = int(year) + (2000 if len(year) == 2 else 0)
    try:
        return date(y, month, int(day))
    except ValueError:
        return None


@dataclass
class ImportProfile:
    header_row: int = 0
    date_col: str = ""
    date_format: str = "auto"
    amount_mode: str = "single"  # single | split
    amount_col: str = ""
    in_col: str = ""
    out_col: str = ""
    desc_cols: list[str] = field(default_factory=list)
    counterparty_col: str = ""
    id_col: str = ""
    balance_col: str = ""
    status_col: str = ""
    status_ok: str = ""  # es. COMPLETED (Revolut): righe con altro stato vengono saltate
    decimal_sep: str = "auto"
    invert_sign: bool = False

    def to_json(self) -> str:
        return json.dumps(asdict(self))

    @classmethod
    def from_json(cls, raw: str) -> "ImportProfile | None":
        if not raw:
            return None
        data = json.loads(raw)
        return cls(**{k: v for k, v in data.items() if k in cls.__dataclass_fields__})


@dataclass
class ParsedTx:
    date: date
    amount_cents: int
    description: str
    counterparty: str
    external_id: str
    balance_cents: int | None
    row_index: int


# ---------- lettura file ----------

class ExcelNumber(str):
    """Numero letto da una cella Excel: il valore è esatto, il separatore decimale non va indovinato."""

    is_number = True


def _cell_to_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S") if (v.hour or v.minute or v.second) else v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        return ExcelNumber(repr(v) if isinstance(v, float) and v != int(v) else str(int(v)))
    return str(v).strip()


def detect_delimiter(text: str) -> str:
    """Sceglie il separatore che produce lo stesso numero di colonne sul maggior numero di righe.
    Contare i caratteri non basta: nei CSV italiani le virgole decimali superano spesso i ';'."""
    from collections import Counter

    lines = [ln for ln in text.splitlines()[:60] if ln.strip()]
    best, best_score = ";", (-1, -1)
    for d in (";", ",", "\t", "|"):
        widths = [len(r) for r in csv.reader(lines, delimiter=d)]
        common = Counter(w for w in widths if w > 1).most_common(1)
        if not common:
            continue
        width, rows = common[0]
        score = (rows, width)
        if score > best_score:
            best, best_score = d, score
    return best


class UnreadableFile(ValueError):
    pass


TABLE_EXTENSIONS = (".csv", ".txt", ".xlsx", ".xlsm", ".xls", ".pdf")


_HEADER_HINT = re.compile(r"\b(data|date|datum|buchungstag)\b", re.I)


def _starts_with_date(text: str) -> bool:
    return bool(text) and parse_date_value(text) is not None


def _pdf_lines(page) -> list[list[dict]]:
    """Parole della pagina raggruppate per riga (stessa altezza), ordinate da sinistra a destra."""
    words = sorted(page.extract_words(keep_blank_chars=False, x_tolerance=1.5), key=lambda w: (round(w["top"]), w["x0"]))
    lines: list[list[dict]] = []
    bounds: list[list[float]] = []  # [top, bottom] di ogni riga
    for w in words:
        if lines:
            top, bottom = bounds[-1]
            overlap = min(bottom, w["bottom"]) - max(top, w["top"])
            height = min(bottom - top, w["bottom"] - w["top"]) or 1
            # stessa riga se le parole si sovrappongono in verticale per oltre metà altezza
            # (i caratteri in grassetto o di corpo diverso hanno "top" leggermente diversi)
            if overlap / height > 0.5:
                lines[-1].append(w)  # i limiti restano quelli della prima parola: la riga non "cresce"
                continue
        lines.append([w])
        bounds.append([w["top"], w["bottom"]])
    return [sorted(ln, key=lambda w: w["x0"]) for ln in lines]


def _header_columns(line: list[dict]) -> list[tuple[float, float, str]]:
    """Unisce le parole vicine dell'intestazione in colonne: (x0, x1, nome)."""
    cols: list[list] = []
    for w in line:
        if cols and w["x0"] - cols[-1][1] < 6:
            cols[-1][1], cols[-1][2] = w["x1"], cols[-1][2] + " " + w["text"]
        else:
            cols.append([w["x0"], w["x1"], w["text"]])
    return [tuple(c) for c in cols]


_AMOUNT_TOKEN = re.compile(r"^[(+\-\u2212\u2013]?\s*(?:€|EUR)?\s*[\d.,' ]+\d\s*(?:€|EUR)?\)?-?$", re.I)


def _looks_like_amount(text: str) -> bool:
    return bool(_AMOUNT_TOKEN.match(text.strip())) and not _starts_with_date(text)


def _assign(line: list[dict], cols: list[tuple[float, float, str]]) -> list[str]:
    cells = [""] * len(cols)
    for w in line:
        center = (w["x0"] + w["x1"]) / 2
        inside = [i for i, (x0, x1, _) in enumerate(cols) if x0 - 4 <= center <= x1 + 4]
        if inside:
            i = inside[0]
        elif w["x1"] <= cols[0][0]:
            i = 0
        else:
            # testo lungo che sborda: va alla colonna che inizia più vicino alla sua sinistra; numeri al centro più vicino
            numeric = _looks_like_amount(w["text"]) or w["text"].upper() in ("€", "EUR", "-", "+", "\u2212")
            if numeric:
                i = min(range(len(cols)), key=lambda k: abs((cols[k][0] + cols[k][1]) / 2 - center))
            else:
                i = max((k for k in range(len(cols)) if cols[k][0] <= w["x0"] + 4), default=0)
        cells[i] = (cells[i] + " " + w["text"]).strip()
    return cells


def _pdf_by_position(pdf) -> list[list[str]]:
    """Ricostruisce la tabella dalle posizioni delle parole, usando le colonne dell'intestazione."""
    out: list[list[str]] = []
    cols = None
    for page in pdf.pages:
        last_top = None
        for line in _pdf_lines(page):
            text = " ".join(w["text"] for w in line)
            if len(line) >= 3 and _HEADER_HINT.search(text) and not _starts_with_date(text):
                new_cols = _header_columns(line)
                if len(new_cols) >= 3:
                    if cols is None:
                        out.append([c[2] for c in new_cols])
                    cols = new_cols
                    continue
            if cols is None:
                continue
            cells = _assign(line, cols)
            top = line[0]["top"]
            if _starts_with_date(cells[0]) or (len(cells) > 1 and _starts_with_date(cells[1])) or len(out) == 1:
                out.append(cells)
                last_top = top
            elif last_top is not None and 0 < top - last_top <= 24:
                # riga di continuazione (descrizione su più righe), vicina alla precedente:
                # accodo solo testo, mai nella colonna della data; piè di pagina e note lontane vengono ignorati
                # e importi stampati un po' più in basso della data (centrati su righe alte): riempiono le celle vuote
                prev = out[-1]
                for i, c in enumerate(cells[1:], start=1):
                    if not c:
                        continue
                    if _looks_like_amount(c):
                        if not prev[i]:
                            prev[i] = c
                    else:
                        prev[i] = (prev[i] + " " + c).strip()
                last_top = top
    return out


def _date_rows(rows: list[list[str]]) -> int:
    return sum(1 for r in rows if any(_starts_with_date(c) for c in r[:2]))


def _read_pdf(content: bytes) -> list[list[str]]:
    """Estratto conto in PDF: tabelle con bordi, poi colonne ricostruite dalle posizioni, poi testo grezzo."""
    import pdfplumber

    def clean(row) -> list[str]:
        return [re.sub(r"\s+", " ", c or "").strip() for c in row]

    try:
        pdf = pdfplumber.open(io.BytesIO(content))
    except Exception as exc:
        raise UnreadableFile(f"PDF non leggibile ({exc})") from exc
    with pdf:
        ruled: list[list[str]] = []
        for page in pdf.pages:
            for table in page.extract_tables() or []:
                ruled += [clean(r) for r in table if any(c for c in r)]
        if ruled:
            width = max(len(r) for r in ruled)
            ruled = [r + [""] * (width - len(r)) for r in ruled]
        positional = _pdf_by_position(pdf)
        # Alcune banche mettono nei riquadri solo intestazione e piè di pagina, e i movimenti come testo libero:
        # tengo il metodo che trova più righe con una data
        best = max((ruled, positional), key=_date_rows)
        if _date_rows(best) > 0:
            return best
        lines = []
        for page in pdf.pages:
            text = page.extract_text(layout=True) or ""
            lines += [re.split(r"\s{2,}", ln.strip()) for ln in text.splitlines() if ln.strip()]
    if not lines:
        raise UnreadableFile("il PDF non contiene testo (forse è una scansione): scarica il CSV o l'Excel dall'home banking")
    width = max(len(r) for r in lines)
    return [r + [""] * (width - len(r)) for r in lines]


_TABLE_CACHE: dict[str, list[list[str]]] = {}


def read_table(content: bytes, filename: str) -> list[list[str]]:
    """Legge CSV, Excel o PDF e restituisce le righe come liste di stringhe. Solleva UnreadableFile."""
    key = hashlib.sha256(content).hexdigest() + Path(filename).suffix.lower()
    if key in _TABLE_CACHE:
        return [list(r) for r in _TABLE_CACHE[key]]
    try:
        rows = _read_table(content, filename)
        if len(_TABLE_CACHE) >= 8:
            _TABLE_CACHE.pop(next(iter(_TABLE_CACHE)))
        _TABLE_CACHE[key] = rows
        return [list(r) for r in rows]
    except UnreadableFile:
        raise
    except Exception as exc:
        raise UnreadableFile(f"file non leggibile ({exc.__class__.__name__}: {exc})") from exc


def _read_table(content: bytes, filename: str) -> list[list[str]]:
    name = filename.lower()
    if name.endswith(".pdf") or content[:5] == b"%PDF-":
        return _read_pdf(content)
    if name.endswith((".xlsx", ".xlsm")):
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [[_cell_to_str(c) for c in row] for row in ws.iter_rows(values_only=True)]
    if name.endswith(".xls"):
        import xlrd

        book = xlrd.open_workbook(file_contents=content)
        sh = book.sheet_by_index(0)
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(_cell_to_str(xlrd.xldate_as_datetime(cell.value, book.datemode)))
                else:
                    row.append(_cell_to_str(cell.value))
            rows.append(row)
        return rows
    if b"\x00" in content[:4096]:
        raise UnreadableFile("non è un file di testo CSV: usa CSV, Excel o PDF")
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    delimiter = detect_delimiter(text)
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


# ---------- euristiche ----------

_KEYWORDS = {
    "date_col": ["data contabile", "data operazione", "date completed", "completed date", "data valuta", "data", "date"],
    "amount_col": ["importo", "total amount", "amount", "importo (eur)", "importo eur"],
    "in_col": ["avere", "entrate", "accrediti", "money in", "credit", "soldi in entrata", "in entrata", "entrata", "eingang", "gutschrift"],
    "out_col": ["dare", "uscite", "addebiti", "money out", "debit", "soldi in uscita", "in uscita", "uscita", "ausgang", "lastschrift"],
    "desc_cols": ["descrizione", "causale", "description", "reference", "dettagli", "descrizione operazione"],
    "counterparty_col": ["controparte", "beneficiario", "payer", "payee", "merchant", "ordinante", "counterparty"],
    "id_col": ["id", "transaction id", "id transazione"],
    "balance_col": ["saldo", "balance"],
    "status_col": ["state", "stato", "status"],
}


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", s.strip().lower())


def detect_header_row(rows: list[list[str]]) -> int:
    width = max((sum(1 for c in r if c) for r in rows[:40]), default=0)
    for i, row in enumerate(rows[:40]):
        filled = [c for c in row if c]
        if len(filled) >= max(2, int(width * 0.6)) and any(re.search(r"\bdata\b|\bdate\b", _norm(c)) for c in filled):
            return i
    for i, row in enumerate(rows[:40]):
        if sum(1 for c in row if c) == width:
            return i
    return 0


def _find(header: list[str], keys: list[str], exclude: set[str] = frozenset()) -> str:
    normed = [(_norm(h), h) for h in header if h and h not in exclude]
    for k in keys:  # prima corrispondenza esatta, nell'ordine di priorità
        for n, h in normed:
            if n == k:
                return h
    for k in keys:
        for n, h in normed:
            if k in n and len(k) > 2:
                return h
    return ""


def guess_profile(rows: list[list[str]]) -> ImportProfile:
    hr = detect_header_row(rows)
    header = rows[hr] if rows else []
    p = ImportProfile(header_row=hr)
    p.date_col = _find(header, _KEYWORDS["date_col"])
    used = {p.date_col}
    p.in_col = _find(header, _KEYWORDS["in_col"], used)
    p.out_col = _find(header, _KEYWORDS["out_col"], used | {p.in_col})
    p.amount_col = _find(header, _KEYWORDS["amount_col"], used | {p.in_col, p.out_col})
    if p.in_col and p.out_col and not p.amount_col:
        p.amount_mode = "split"
    used |= {p.amount_col, p.in_col, p.out_col}
    desc = _find(header, _KEYWORDS["desc_cols"], used)
    p.desc_cols = [desc] if desc else []
    used |= set(p.desc_cols)
    p.counterparty_col = _find(header, _KEYWORDS["counterparty_col"], used)
    p.id_col = next((h for h in header if _norm(h) in _KEYWORDS["id_col"]), "")
    p.balance_col = _find(header, _KEYWORDS["balance_col"], used)
    p.status_col = next((h for h in header if _norm(h) in _KEYWORDS["status_col"]), "")
    if p.status_col:
        values = {r[header.index(p.status_col)].upper() for r in rows[hr + 1 : hr + 200] if len(r) > header.index(p.status_col)}
        p.status_ok = "COMPLETED" if "COMPLETED" in values else ""
    return p


def parse_date_value(value: str, fmt: str = "auto") -> date | None:
    value = value.strip()
    if not value:
        return None
    formats = DATE_FORMATS if fmt == "auto" else [fmt]
    for f in formats:
        try:
            return datetime.strptime(value, f).date()
        except ValueError:
            continue
    if fmt != "auto":
        return _text_date(value)
    # Excel a volte esporta il numero seriale
    if re.fullmatch(r"\d{5}(\.\d+)?", value):
        return date(1899, 12, 30) + timedelta(days=int(float(value)))
    # mese a lettere, o data seguita da altro testo (es. "01/10/2026 14:32 CET", "01.10.2026 Valuta")
    d = _text_date(value)
    if d:
        return d
    m = _LEADING_DATE.match(value)
    if m and m.group(1) != value:
        return parse_date_value(m.group(1))
    return None


def parse_rows(rows: list[list[str]], p: ImportProfile) -> tuple[list[ParsedTx], list[str]]:
    header = rows[p.header_row] if rows else []
    idx = {h: i for i, h in enumerate(header) if h}

    def get(row: list[str], col: str) -> str:
        i = idx.get(col)
        return row[i] if i is not None and i < len(row) else ""

    dec = None if p.decimal_sep == "auto" else p.decimal_sep
    out: list[ParsedTx] = []
    errors: list[str] = []
    for n, row in enumerate(rows[p.header_row + 1 :], start=p.header_row + 2):
        if not any(row):
            continue
        if p.status_col and p.status_ok and get(row, p.status_col).upper() != p.status_ok.upper():
            continue
        raw_date = get(row, p.date_col)
        d = parse_date_value(raw_date, p.date_format)
        if d is None:
            # righe di totale/piè di pagina senza importo: le salto senza segnalarle
            amount_cols = (p.in_col, p.out_col) if p.amount_mode == "split" else (p.amount_col,)
            if any(parse_amount(get(row, c), dec) for c in amount_cols if c):
                errors.append(f"Riga {n}: data non riconosciuta '{raw_date}'")
            continue
        if p.amount_mode == "split":
            a_in = parse_amount(get(row, p.in_col), dec) or 0
            a_out = parse_amount(get(row, p.out_col), dec) or 0
            amount = abs(a_in) - abs(a_out)
        else:
            raw_amount = get(row, p.amount_col)
            amount = parse_amount(raw_amount, dec)
            if amount is None:
                shown = f"'{raw_amount[:40]}'" if raw_amount else "cella vuota"
                errors.append(f"Riga {n}: importo non riconosciuto ({shown})")
                continue
        if p.invert_sign:
            amount = -amount
        desc = " · ".join(v for v in (get(row, c) for c in p.desc_cols) if v)
        bal = parse_amount(get(row, p.balance_col), dec) if p.balance_col else None
        out.append(
            ParsedTx(
                date=d,
                amount_cents=amount,
                description=desc,
                counterparty=get(row, p.counterparty_col) if p.counterparty_col else "",
                external_id=get(row, p.id_col) if p.id_col else "",
                balance_cents=bal,
                row_index=n,
            )
        )
    if p.amount_mode == "single" and not p.invert_sign:
        _signs_from_balance(out)
    return out, errors


def _signs_from_balance(txs: list[ParsedTx]) -> None:
    """Se gli importi sono tutti positivi (es. PDF con colonne separate non riconosciute) ma c'è il saldo,
    il segno si ricava dalla variazione del saldo tra un movimento e il successivo."""
    if not txs or any(t.amount_cents < 0 for t in txs):
        return
    with_bal = [t for t in txs if t.balance_cents is not None]
    if len(with_bal) < max(2, len(txs) // 2):
        return
    ordered = sorted(with_bal, key=lambda t: t.row_index)
    descending = ordered[0].date > ordered[-1].date  # estratti con i più recenti in alto
    for prev, cur in zip(ordered, ordered[1:]):
        older, newer = (cur, prev) if descending else (prev, cur)
        delta = newer.balance_cents - older.balance_cents
        if abs(delta) == newer.amount_cents and delta < 0:
            newer.amount_cents = -newer.amount_cents


def new_batch_id(prefix: str = "") -> str:
    """Identificativo univoco di un import (serve ad annullarlo senza toccare altri import)."""
    import uuid

    return f"{prefix}{datetime.now():%Y%m%d%H%M%S}-{uuid.uuid4().hex[:6]}"


def dedupe_keys(txs: list[ParsedTx]) -> list[str]:
    """Chiave stabile per riconoscere un movimento già importato anche se il file viene ricaricato."""
    seen: dict[str, int] = {}
    keys = []
    for t in txs:
        if t.external_id:
            base = f"id|{t.external_id}"
        else:
            base = f"{t.date.isoformat()}|{t.amount_cents}|{_norm(t.description)}"
            occ = seen.get(base, 0)
            seen[base] = occ + 1
            base = f"{base}|{occ}"
        keys.append(hashlib.sha256(base.encode()).hexdigest()[:40])
    return keys


@dataclass
class ImportResult:
    inserted: int
    skipped: int
    snapshot: tuple[date, int] | None


def closing_balance(txs: list[ParsedTx]) -> tuple[date, int] | None:
    """Saldo a fine dell'ultimo giorno, se il file ha una colonna saldo."""
    with_bal = [t for t in txs if t.balance_cents is not None]
    if not with_bal:
        return None
    last_day = max(t.date for t in with_bal)
    same_day = [t for t in with_bal if t.date == last_day]
    # Ordine del file: se è decrescente (più recenti in alto) l'ultimo movimento del giorno è il primo
    descending = with_bal[0].date > with_bal[-1].date
    pick = min(same_day, key=lambda t: t.row_index) if descending else max(same_day, key=lambda t: t.row_index)
    return last_day, pick.balance_cents


def import_transactions(
    session: Session, account: Account, txs: list[ParsedTx], batch: str, update_balance: bool = True, source: str = "import"
) -> ImportResult:
    keys = dedupe_keys(txs)
    existing = set(
        session.scalars(select(Transaction.dedupe_key).where(Transaction.account_id == account.id, Transaction.dedupe_key.in_(keys)))
    )
    inserted = 0
    new_txs: list[Transaction] = []
    for t, k in zip(txs, keys):
        if k in existing:
            continue
        existing.add(k)
        new_txs.append(
            Transaction(
                account_id=account.id,
                date=t.date,
                amount_cents=t.amount_cents,
                description=t.description,
                counterparty=t.counterparty,
                source=source,
                dedupe_key=k,
                import_batch=batch,
            )
        )
        session.add(new_txs[-1])
        inserted += 1
    if new_txs:
        from app.services.categorize import apply as categorize

        categorize(session, new_txs)
    snap = closing_balance(txs) if update_balance else None
    if snap:
        row = session.scalar(select(BalanceSnapshot).where(BalanceSnapshot.account_id == account.id, BalanceSnapshot.date == snap[0]))
        if row:
            row.balance_cents, row.source = snap[1], "import"
        else:
            session.add(BalanceSnapshot(account_id=account.id, date=snap[0], balance_cents=snap[1], source="import"))
    session.commit()
    return ImportResult(inserted=inserted, skipped=len(txs) - inserted, snapshot=snap)
