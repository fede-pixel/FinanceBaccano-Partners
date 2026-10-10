"""Scadenzario dei documenti da consegnare al commercialista."""
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from app import config
from app.models import Account, ChecklistItem, DocCategory, Document, Transaction
from app.services.forecast import calculators
from app.services.util import add_months, fmt_date, get_setting, month_bounds, month_label, set_setting

STATUS_LABELS = {
    "da_fare": "Da fare",
    "completato": "Pronto",
    "non_applicabile": "Non applicabile",
    "inviato": "Inviato",
}

DEFAULT_CATEGORIES = [
    dict(code="fatture_estere", name="Fatture/ricevute acquisti online da fornitori esteri",
         description="Fatture e/o ricevute di tutti gli acquisti (beni e servizi) effettuati online da fornitori con sede fuori dall'Italia.",
         frequency="mensile", due_day=6, due_month_offset=1),
    dict(code="movimenti_banche", name="Lista movimenti banche",
         description="Lista movimenti di tutti i conti correnti in essere.",
         frequency="mensile", due_day=6, due_month_offset=1, auto_export="conto"),
    dict(code="ec_carte", name="Estratti conto carte di credito / prepagate",
         description="Estratti conto di tutte le carte di credito e prepagate in essere.",
         frequency="mensile", due_day=6, due_month_offset=1, auto_export="carta"),
    dict(code="ec_trimestrali", name="Estratti conto ufficiali trimestrali banche",
         description="Estratti conto 'ufficiali' trimestrali. Di solito disponibili nell'home banking intorno al 15 del mese successivo alla chiusura del trimestre.",
         frequency="trimestrale", due_day=15, due_month_offset=1),
    dict(code="contanti", name="Elenco pagamenti e incassi in contanti",
         description="Elenco dei pagamenti e incassi effettuati in contanti (prima nota cassa).",
         frequency="mensile", due_day=6, due_month_offset=1, auto_export="cassa"),
    dict(code="assegni_emessi", name="Copia assegni emessi",
         description="Copia degli assegni bancari emessi per pagamenti.",
         frequency="mensile", due_day=6, due_month_offset=1),
    dict(code="assegni_ricevuti", name="Copia assegni ricevuti e versati",
         description="Copia degli assegni bancari ricevuti e versati in banca.",
         frequency="mensile", due_day=6, due_month_offset=1),
    dict(code="spese_varie", name="Spese senza fattura elettronica",
         description="Copia di tutte le spese diverse da fattura elettronica (autostrada, assicurazioni, ricevute bar/ristoranti, ecc.).",
         frequency="mensile", due_day=6, due_month_offset=1),
]


def seed_categories(session: Session, today: date | None = None) -> None:
    today = today or date.today()
    existing = {c.code for c in session.scalars(select(DocCategory))}
    for order, data in enumerate(DEFAULT_CATEGORIES):
        if data["code"] not in existing:
            session.add(DocCategory(sort_order=order, **data))
    if not get_setting(session, "admin_start_period"):
        prev = add_months(today.replace(day=1), -1)
        set_setting(session, "admin_start_period", f"{prev.year}-{prev.month:02d}")
    session.commit()


# ---------- periodi ----------

def month_key(d: date) -> str:
    return f"{d.year}-{d.month:02d}"


def quarter_key(d: date) -> str:
    return f"{d.year}-Q{(d.month - 1) // 3 + 1}"


def period_label(key: str) -> str:
    if "-Q" in key:
        y, q = key.split("-Q")
        return f"{q}° trimestre {y}"
    y, m = key.split("-")
    return month_label(int(y), int(m))


def period_range(key: str) -> tuple[date, date]:
    if "-Q" in key:
        y, q = key.split("-Q")
        first_month = (int(q) - 1) * 3 + 1
        return date(int(y), first_month, 1), month_bounds(int(y), first_month + 2)[1]
    y, m = key.split("-")
    return month_bounds(int(y), int(m))


def due_date_for(cat: DocCategory, period_end: date) -> date:
    return add_months(period_end.replace(day=1), cat.due_month_offset, day=cat.due_day)


def ensure_checklist(session: Session, today: date | None = None) -> None:
    """Crea le righe dello scadenzario per ogni periodo iniziato (dal mese di partenza a oggi)."""
    today = today or date.today()
    start_key = get_setting(session, "admin_start_period") or month_key(today)
    y, m = (int(x) for x in start_key.split("-"))
    start = date(y, m, 1)
    cats = list(session.scalars(select(DocCategory).where(DocCategory.active)))
    have = {(c, k) for c, k in session.execute(select(ChecklistItem.category_id, ChecklistItem.period_key))}
    d = start
    while d <= today:
        for cat in cats:
            key = quarter_key(d) if cat.frequency == "trimestrale" else month_key(d)
            if (cat.id, key) in have:
                continue
            _, end = period_range(key)
            session.add(ChecklistItem(category_id=cat.id, period_key=key, period_end=end, due_date=due_date_for(cat, end)))
            have.add((cat.id, key))
        d = add_months(d, 1, day=1)
    session.commit()


@dataclass
class DeadlineInfo:
    item: ChecklistItem
    days_left: int
    overdue: bool


def open_deadlines(session: Session, today: date, within_days: int = 45) -> list[DeadlineInfo]:
    items = session.scalars(
        select(ChecklistItem).where(ChecklistItem.status == "da_fare").order_by(ChecklistItem.due_date)
    )
    out = []
    for it in items:
        left = (it.due_date - today).days
        if left <= within_days:
            out.append(DeadlineInfo(it, left, left < 0))
    return out


# ---------- archivio file ----------

def docs_root(session: Session) -> Path:
    p = get_setting(session, "docs_dir")
    root = Path(p).expanduser() if p else config.DEFAULT_DOCS_DIR
    root.mkdir(parents=True, exist_ok=True)
    return root


def packages_root(session: Session) -> Path:
    p = get_setting(session, "packages_dir")
    root = Path(p).expanduser() if p else config.DEFAULT_PACKAGES_DIR
    root.mkdir(parents=True, exist_ok=True)
    return root


def _safe(name: str) -> str:
    name = re.sub(r"[\\/:*?\"<>|\x00-\x1f]", "_", name).strip(" .")
    return name or "documento"


def item_folder(item: ChecklistItem) -> Path:
    year = item.period_key[:4]
    return Path(year) / item.period_key / f"{item.category.sort_order + 1:02d} {_safe(item.category.name)}"


def store_document(session: Session, item: ChecklistItem, filename: str, content: bytes, **meta) -> Document:
    rel_dir = item_folder(item)
    target_dir = docs_root(session) / rel_dir
    target_dir.mkdir(parents=True, exist_ok=True)
    base = _safe(Path(filename).name)
    stem, suffix = Path(base).stem, Path(base).suffix
    target = target_dir / base
    n = 2
    while target.exists():
        target = target_dir / f"{stem} ({n}){suffix}"
        n += 1
    target.write_bytes(content)
    doc = Document(
        checklist_item=item,
        original_name=filename,
        stored_path=(rel_dir / target.name).as_posix(),
        size=len(content),
        **meta,
    )
    session.add(doc)
    if item.status == "da_fare":
        item.status = "completato"
        item.completed_at = datetime.now()
    return doc


def document_path(session: Session, doc: Document) -> Path:
    return docs_root(session) / doc.stored_path


def delete_document(session: Session, doc: Document) -> None:
    path = document_path(session, doc)
    if path.exists():
        path.unlink()
    item = doc.checklist_item
    item.documents.remove(doc)  # delete-orphan: viene eliminato dal DB
    session.flush()
    if item.status == "completato" and not item.documents:
        item.status = "da_fare"
        item.completed_at = None


# ---------- export automatici dai movimenti ----------

def export_movements_xlsx(session: Session, account: Account, start: date, end: date) -> bytes:
    import io

    import openpyxl
    from openpyxl.styles import Font, PatternFill

    calc = calculators(session, [account])[account.id]
    opening = calc.at(date.fromordinal(start.toordinal() - 1))
    txs = list(
        session.scalars(
            select(Transaction)
            .where(Transaction.account_id == account.id, Transaction.date >= start, Transaction.date <= end)
            .order_by(Transaction.date, Transaction.id)
        )
    )
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Movimenti"
    bold = Font(bold=True)
    ws.append([f"{account.name} – {account.bank}".strip(" –")])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Periodo: {fmt_date(start)} – {fmt_date(end)}" + (f" · IBAN {account.iban}" if account.iban else "")])
    ws.append([])
    ws.append(["Data", "Descrizione", "Controparte", "Categoria", "Entrate", "Uscite", "Saldo"])
    for c in ws[4]:
        c.font = bold
        c.fill = PatternFill("solid", fgColor="E1E0D9")
    ws.append([start, "Saldo iniziale", "", "", None, None, opening / 100])
    running = opening
    tot_in = tot_out = 0
    for t in txs:
        running += t.amount_cents
        if t.amount_cents >= 0:
            tot_in += t.amount_cents
        else:
            tot_out += t.amount_cents
        ws.append([
            t.date, t.description, t.counterparty, t.category,
            t.amount_cents / 100 if t.amount_cents >= 0 else None,
            -t.amount_cents / 100 if t.amount_cents < 0 else None,
            running / 100,
        ])
    ws.append([end, "Totali e saldo finale", "", "", tot_in / 100, -tot_out / 100, running / 100])
    for c in ws[ws.max_row]:
        c.font = bold
    for row in ws.iter_rows(min_row=5):
        row[0].number_format = "DD/MM/YYYY"
        for c in row[4:7]:
            c.number_format = '#,##0.00 "€"'
    for col, width in zip("ABCDEFG", (12, 60, 30, 18, 14, 14, 14)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_auto_exports(session: Session, item: ChecklistItem) -> list[Document]:
    """Per le categorie con export automatico crea un Excel per ogni conto/carta/cassa del tipo giusto."""
    kind = item.category.auto_export
    if not kind:
        return []
    start, end = period_range(item.period_key)
    accounts = list(session.scalars(select(Account).where(Account.kind == kind, Account.active)))
    # Rigenera: tolgo gli export precedenti dello stesso periodo
    for doc in [d for d in item.documents if d.generated]:
        delete_document(session, doc)
    docs = []
    for acc in accounts:
        content = export_movements_xlsx(session, acc, start, end)
        label = "Prima nota cassa" if kind == "cassa" else "Movimenti"
        name = f"{label} {acc.name} {item.period_key}.xlsx"
        docs.append(store_document(session, item, name, content, generated=True, notes="Generato dai movimenti importati"))
    session.commit()
    return docs


# ---------- pacchetto per il commercialista ----------

def package_items(session: Session, key: str) -> list[ChecklistItem]:
    """Elementi da includere nel pacchetto del mese `key` (es. 2026-09): tutti i mensili di quel mese
    più i trimestrali chiusi entro quel mese e non ancora inviati."""
    _, month_end = period_range(key)
    monthly = select(ChecklistItem).join(DocCategory).where(ChecklistItem.period_key == key)
    quarterly = (
        select(ChecklistItem)
        .join(DocCategory)
        .where(DocCategory.frequency == "trimestrale", ChecklistItem.period_end <= month_end, ChecklistItem.status != "inviato")
    )
    items = {i.id: i for i in session.scalars(monthly)}
    for i in session.scalars(quarterly):
        if i.documents:
            items[i.id] = i
    return sorted(items.values(), key=lambda i: (i.category.sort_order, i.period_key))


def build_package(session: Session, key: str) -> Path:
    items = package_items(session, key)
    company = _safe(get_setting(session, "company_name") or "Documenti")
    out_dir = packages_root(session)
    zip_path = out_dir / f"{company} - Documenti {key}.zip"
    root = docs_root(session)
    lines = [
        f"{get_setting(session, 'company_name')} – documenti {period_label(key)}",
        f"Generato il {fmt_date(date.today())}",
        "",
    ]
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for it in items:
            folder = f"{it.category.sort_order + 1:02d} {_safe(it.category.name)}"
            if it.category.frequency == "trimestrale":
                folder += f" ({it.period_key})"
            status = STATUS_LABELS.get(it.status, it.status)
            lines.append(f"[{status}] {it.category.name} – {period_label(it.period_key)}")
            if it.notes:
                lines.append(f"    Note: {it.notes}")
            for doc in it.documents:
                src = root / doc.stored_path
                if src.exists():
                    zf.write(src, f"{folder}/{src.name}")
                    lines.append(f"    - {src.name}")
            if not it.documents:
                lines.append("    (nessun documento)" if it.status != "non_applicabile" else "    (nessuna operazione nel periodo)")
            lines.append("")
        zf.writestr("Riepilogo.txt", "\n".join(lines))
    return zip_path


def mark_sent(session: Session, items: list[ChecklistItem]) -> None:
    now = datetime.now()
    for it in items:
        if it.status != "inviato" and (it.documents or it.status in ("completato", "non_applicabile")):
            it.status = "inviato"
            it.sent_at = now
    session.commit()

